from datetime import date, timedelta
import csv
import io
import json
import hashlib
import secrets
from django.contrib.auth import login, logout
from django.contrib.sessions.models import Session
from django.core.mail import send_mail
from django.contrib.auth.password_validation import validate_password
from django.conf import settings
from django.views.decorators.csrf import ensure_csrf_cookie
from django.db.models.deletion import ProtectedError
from django.db import transaction
from django.http import HttpResponse
from django.db.models import Q
from django.utils import timezone
from django.core.validators import validate_email
from django.core.paginator import Paginator
from django.core.exceptions import ValidationError as DjangoValidationError
from rest_framework import status, viewsets
from rest_framework.decorators import action, api_view, permission_classes
from rest_framework.permissions import AllowAny
from rest_framework.response import Response
from rest_framework.exceptions import PermissionDenied, ValidationError
from .models import Activity, Contact, CrmSession, Deal, ImportBatch, Lead, Notification, PasswordResetToken, Permission, Product, Record, Role, Source, Task, User, LearningMode, CourseDuration
from .serializers import ActivitySerializer, ContactSerializer, DealSerializer, HistorySerializer, LeadSerializer, NotificationSerializer, ProductSerializer, RoleSerializer, SourceSerializer, TaskSerializer, UserSummarySerializer, LearningModeSerializer, CourseDurationSerializer
from .services import assert_visible, audit, can_read_record_history, convert_lead, history, record_is_visible
from .authentication import find_user, normalize_identifier


def visible_records(user):
    query = Record.objects.all()
    return query if user.has_crm_permission("view_workspace_records") else query.filter(owner=user)


def require_permission(request, key):
    if not request.user.has_crm_permission(key):
        raise PermissionDenied("You do not have permission to perform this action.")


def _history_changes(before, after, fields):
    changed = {}
    for field in fields:
        old = before.get(field)
        new = after.get(field)
        if old != new:
            changed[field] = {"from": old, "to": new}
    return changed


class LeadViewSet(viewsets.ModelViewSet):
    serializer_class = LeadSerializer
    permission_key = "view_leads"

    def get_queryset(self):
        query = Lead.objects.filter(
            record__in=visible_records(self.request.user),
            deleted_at__isnull=True,
        ).select_related("record__owner", "source", "product", "learning_mode", "course_duration")
        search = self.request.query_params.get("search", "").strip()
        if search:
            query = query.filter(Q(name__icontains=search) | Q(company__icontains=search) | Q(email__icontains=search))
        status_value = self.request.query_params.get("status")
        if status_value:
            query = query.filter(status=status_value)
        source_value = self.request.query_params.get("source")
        if source_value:
            query = query.filter(source_id=source_value)
        owner_value = self.request.query_params.get("owner")
        if owner_value:
            query = query.filter(record__owner_id=owner_value)
        return query.order_by("-created_at", "-id")

    def list(self, request, *args, **kwargs):
        require_permission(request, "view_leads")
        return super().list(request, *args, **kwargs)

    def retrieve(self, request, *args, **kwargs):
        require_permission(request, "view_leads")
        return super().retrieve(request, *args, **kwargs)

    def create(self, request, *args, **kwargs):
        require_permission(request, "capture_lead")
        payload = request.data.copy()
        owner_id = payload.pop("owner", None) or request.user.id
        try:
            owner_id = int(owner_id)
        except (TypeError, ValueError):
            return Response({"owner": "Invalid owner."}, status=400)
        owner = User.objects.filter(pk=owner_id, is_active=True).first()
        if not owner:
            return Response({"owner": "Owner must be an active user."}, status=400)
        if owner.id != request.user.id and not request.user.has_crm_permission("reassign_lead"):
            return Response({"detail": "You do not have permission to assign this lead to another user."}, status=403)
        if payload.get("status") == Lead.Status.QUALIFIED and not request.user.has_crm_permission("set_lead_status"):
            return Response(status=403)
        if not payload.get("source"):
            fallback_source = Source.objects.filter(name="Other", active=True).first()
            if fallback_source:
                payload["source"] = fallback_source.id
        serializer = self.get_serializer(data=payload)
        serializer.is_valid(raise_exception=True)
        email = (serializer.validated_data.get("email") or "").strip()
        duplicate_warning = (
            Lead.objects.filter(email__iexact=email, deleted_at__isnull=True).exists()
            if email
            else False
        )
        lead = serializer.save(record=Record.objects.create(owner=owner))
        history(
            lead.record,
            request.user,
            "lead_captured",
            {
                "source": lead.source.name if lead.source else None,
                "owner": owner.get_full_name() or owner.username,
            },
        )
        audit(request.user, "lead.created", lead, after=serializer.data)
        response_data = self.get_serializer(lead).data
        response_data["duplicate_warning"] = duplicate_warning
        return Response(response_data, status=status.HTTP_201_CREATED)

    def update(self, request, *args, **kwargs):
        require_permission(request, "edit_lead")
        return self._update(request, *args, **kwargs)

    def partial_update(self, request, *args, **kwargs):
        require_permission(request, "edit_lead")
        return self._update(request, *args, partial=True, **kwargs)

    def _update(self, request, *args, partial=False, **kwargs):
        lead = self.get_object()
        payload = request.data.copy()
        has_owner_change = "owner" in payload
        owner_value = payload.pop("owner", None) if has_owner_change else None
        if owner_value is not None:
            try:
                owner_id = int(owner_value)
            except (TypeError, ValueError):
                return Response({"owner": "Invalid owner."}, status=400)
            if not request.user.has_crm_permission("reassign_lead") and owner_id != lead.record.owner_id:
                return Response(status=403)
            owner = User.objects.filter(pk=owner_id, is_active=True).first()
            if not owner:
                return Response({"owner": "Owner must be an active user."}, status=400)
        if "status" in payload and payload.get("status") != lead.status:
            if not request.user.has_crm_permission("set_lead_status"):
                return Response(status=403)
            if payload.get("status") not in Lead.Status.values:
                return Response({"status": "Invalid lead status."}, status=400)
            # Disqualification reason is optional. Validate it only when supplied.
            if payload.get("status") == Lead.Status.NOT_QUALIFIED and payload.get("disqualification_reason"):
                reason = str(payload.get("disqualification_reason") or "").strip()
                allowed = {"no budget", "no need", "too small", "wrong fit", "went elsewhere", "no response", "other"}
                if reason not in allowed:
                    return Response({"disqualification_reason": "Choose a valid disqualification reason."}, status=400)
                if reason == "other" and payload.get("notes") and len((payload.get("notes") or "").strip()) < 5:
                    return Response({"notes": "Please provide a sentence explaining the reason."}, status=400)
        before = self.get_serializer(lead).data
        serializer = self.get_serializer(lead, data=payload, partial=partial)
        serializer.is_valid(raise_exception=True)
        updated = serializer.save()
        if has_owner_change:
            old_owner = lead.record.owner_id
            new_owner = int(owner_value)
            if old_owner != new_owner:
                lead.record.owner_id = new_owner
                lead.record.save(update_fields=["owner"])
                old_owner_obj = User.objects.filter(pk=old_owner).first()
                new_owner_obj = User.objects.filter(pk=new_owner).first()
                history(
                    lead.record,
                    request.user,
                    "reassigned",
                    {
                        "from_owner": old_owner_obj.get_full_name() or old_owner_obj.username if old_owner_obj else str(old_owner),
                        "to_owner": new_owner_obj.get_full_name() or new_owner_obj.username if new_owner_obj else str(new_owner),
                    },
                )
                audit(request.user, "lead.reassigned", lead, before={"owner_id": old_owner}, after={"owner_id": new_owner})
        if "status" in payload and payload.get("status") != before.get("status"):
            history(
                lead.record,
                request.user,
                "qualified" if updated.status == Lead.Status.QUALIFIED else "not_qualified",
                {"from": before.get("status"), "to": updated.status, "reason": updated.disqualification_reason},
            )
        after = self.get_serializer(updated).data
        changed = _history_changes(
            before,
            after,
            ("name", "company", "email", "phone", "source", "product", "learning_mode", "course_duration", "town_city", "notes"),
        )
        if changed:
            audit(request.user, "lead.updated", updated, before=before, after=after)
            history(lead.record, request.user, "edited", {"changed": changed})
        return Response(self.get_serializer(updated).data)

    def destroy(self, request, *args, **kwargs):
        require_permission(request, "delete_lead")
        lead = self.get_object()
        lead.deleted_at = timezone.now()
        lead.save(update_fields=["deleted_at"])
        history(lead.record, request.user, "deleted", {"entity": "Lead", "name": lead.name, "company": lead.company})
        audit(request.user, "lead.deleted", lead, after={"deleted_at": lead.deleted_at.isoformat()})
        return Response(status=204)

    @action(detail=True, methods=["post"])
    def qualify(self, request, pk=None):
        require_permission(request, "set_lead_status")
        lead = self.get_object()
        new_status = request.data.get("status")
        reason = (request.data.get("reason") or "").strip()
        if new_status not in Lead.Status.values:
            return Response({"status": "Invalid lead status."}, status=400)
        if new_status == Lead.Status.NOT_QUALIFIED and reason:
            allowed = {"no budget", "no need", "too small", "wrong fit", "went elsewhere", "no response", "other"}
            if reason not in allowed:
                return Response({"reason": "Choose a valid disqualification reason."}, status=400)
            if reason == "other" and request.data.get("details") and len((request.data.get("details") or "").strip()) < 5:
                return Response({"details": "Please provide a sentence explaining the reason."}, status=400)
        old = lead.status
        lead.status = new_status
        lead.disqualification_reason = reason if new_status == Lead.Status.NOT_QUALIFIED else ""
        lead.save(update_fields=["status", "disqualification_reason"])
        history(
            lead.record,
            request.user,
            "qualified" if new_status == Lead.Status.QUALIFIED else "not_qualified",
            {"from": old, "to": new_status, "reason": lead.disqualification_reason},
        )
        audit(request.user, "lead.status_changed", lead, before={"status": old}, after={"status": new_status, "reason": lead.disqualification_reason})
        return Response(self.get_serializer(lead).data)

    @action(detail=True, methods=["post"])
    def reassign(self, request, pk=None):
        require_permission(request, "reassign_lead")
        lead = self.get_object()
        owner = User.objects.filter(pk=request.data.get("owner"), is_active=True).first()
        if not owner:
            return Response({"owner": "Owner must be an active user."}, status=400)
        old_owner = lead.record.owner_id
        if old_owner == owner.id:
            return Response(self.get_serializer(lead).data)
        lead.record.owner = owner
        lead.record.save(update_fields=["owner"])
        history(
            lead.record,
            request.user,
            "reassigned",
            {
                "from_owner": (User.objects.filter(pk=old_owner).first().get_full_name() if User.objects.filter(pk=old_owner).first() else str(old_owner)),
                "to_owner": (owner.get_full_name() or owner.username),
            },
        )
        audit(request.user, "lead.reassigned", lead, before={"owner_id": old_owner}, after={"owner_id": owner.id})
        return Response(self.get_serializer(lead).data)

    @action(detail=True, methods=["post"])
    def convert(self, request, pk=None):
        require_permission(request, "move_to_contact")
        lead = self.get_object()
        contact, activity_count = convert_lead(lead, request.user)
        return Response({"contact": ContactSerializer(contact).data, "activity_count": activity_count}, status=201)

    @action(detail=False, methods=["get"], url_path="import-template")
    def import_template(self, request):
        """Download a workbook documenting every import column and accepted values."""
        require_permission(request, "import_leads")
        from openpyxl import Workbook
        from openpyxl.utils import get_column_letter

        columns = [
            ("name", "Name", "Required. Lead name."),
            ("company", "Company", "Optional."),
            ("email", "Email", "Optional. If provided, must be a valid email address."),
            ("phone", "Phone", "Optional. Indian mobile numbers are accepted."),
            ("source", "Source", "Optional. If provided, must exactly match an active Lead Source."),
            ("product", "Course Interested", "Optional. Must match an active Product."),
            ("learning_mode", "Preferred Learning Mode", "Optional. Must match an active Learning Mode."),
            ("course_duration", "Preferred Course Duration", "Optional. Must match an active Course Duration."),
            ("town_city", "Town / City", "Optional."),
            ("status", "Status", "Optional. Accepted values are listed on the Accepted Values sheet."),
            ("notes", "Lead Feedback", "Optional."),
        ]

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Lead Import Template"
        for index, (_, label, _) in enumerate(columns, start=1):
            sheet.cell(row=1, column=index, value=label)
            sheet.column_dimensions[get_column_letter(index)].width = max(18, len(label) + 4)

        values = workbook.create_sheet("Accepted Values")
        values.append(["Field", "Accepted value"])
        values.append(["Source", "Use active Lead Sources from the CRM."])
        for source in Source.objects.filter(active=True).order_by("name"):
            values.append(["Source", source.name])
        values.append(["Course Interested", "Use active Products from the CRM."])
        for product in Product.objects.filter(active=True).order_by("name"):
            values.append(["Course Interested", product.name])
        values.append(["Preferred Learning Mode", "Use active Learning Modes from the CRM."])
        for mode in LearningMode.objects.filter(active=True).order_by("name"):
            values.append(["Preferred Learning Mode", mode.name])
        values.append(["Preferred Course Duration", "Use active Course Durations from the CRM."])
        for duration in CourseDuration.objects.filter(active=True).order_by("name"):
            values.append(["Preferred Course Duration", duration.name])
        for value, label in Lead.Status.choices:
            values.append(["Status", value])
            values.append(["Status", label])
        values.append(["Duplicate action", "skip"])
        values.append(["Duplicate action", "import_anyway"])
        values.append(["Duplicate action", "update"])

        instructions = workbook.create_sheet("Instructions")
        instructions.append(["LiteCRM Bulk Lead Import"])
        instructions.append(["Required fields", "Name"])
        instructions.append(["Duplicate handling", "Duplicates are flagged during preview. Choose skip, import anyway, or update before importing."])
        instructions.append(["Validation", "Rows with errors must be fixed or explicitly skipped before anything is written."])
        instructions.append(["Accepted values", "Use the Accepted Values sheet for active master-data values."])

        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = 'attachment; filename="litecrm-leads-template.xlsx"'
        return response

    @staticmethod
    def _normalise_import_header(value):
        return " ".join(str(value or "").strip().lower().replace("_", " ").replace("-", " ").split())

    @classmethod
    def _import_column_aliases(cls):
        return {
            "name": "name",
            "lead name": "name",
            "company": "company",
            "email": "email",
            "email address": "email",
            "phone": "phone",
            "mobile": "phone",
            "mobile number": "phone",
            "source": "source",
            "lead source": "source",
            "product": "product",
            "course": "product",
            "course interested": "product",
            "learning mode": "learning_mode",
            "preferred learning mode": "learning_mode",
            "course duration": "course_duration",
            "preferred course duration": "course_duration",
            "town city": "town_city",
            "town / city": "town_city",
            "town/city": "town_city",
            "city": "town_city",
            "status": "status",
            "lead feedback": "notes",
            "notes": "notes",
        }

    @classmethod
    def _parse_import_upload(cls, upload):
        filename = str(upload.name or "").lower()
        if filename.endswith(".xlsx"):
            from openpyxl import load_workbook
            workbook = load_workbook(upload, read_only=True, data_only=True)
            sheet = workbook.active
            rows = list(sheet.iter_rows(values_only=True))
            workbook.close()
            if not rows:
                raise ValueError("The file is empty.")
            headers = [str(value or "").strip() for value in rows[0]]
            data = []
            for row in rows[1:]:
                if not any(value not in (None, "") for value in row):
                    continue
                data.append({headers[index]: row[index] if index < len(row) else "" for index in range(len(headers))})
            return headers, data
        if filename.endswith(".csv"):
            text = upload.read().decode("utf-8-sig")
            reader = csv.DictReader(io.StringIO(text))
            headers = [str(value or "").strip() for value in (reader.fieldnames or [])]
            if not headers:
                raise ValueError("The file does not contain a header row.")
            return headers, list(reader)
        raise ValueError("Only CSV and XLSX files are supported.")

    @classmethod
    def _suggest_import_mapping(cls, headers):
        aliases = cls._import_column_aliases()
        mapping = {}
        used = set()
        for header in headers:
            canonical = aliases.get(cls._normalise_import_header(header))
            if canonical and canonical not in mapping and header not in used:
                mapping[canonical] = header
                used.add(header)
        return mapping

    @classmethod
    def _apply_import_mapping(cls, raw_row, mapping):
        if not mapping:
            mapping = cls._suggest_import_mapping(list(raw_row.keys()))
        result = {}
        for canonical, source_header in mapping.items():
            if source_header in raw_row:
                result[canonical] = raw_row.get(source_header)
        # Accept mappings supplied as source-header -> canonical as well.
        if not result:
            aliases = cls._import_column_aliases()
            for source_header, canonical in mapping.items():
                if canonical in aliases:
                    canonical = aliases[canonical]
                if source_header in raw_row:
                    result[canonical] = raw_row.get(source_header)
        return {key: "" if value is None else str(value).strip() for key, value in result.items()}

    @classmethod
    def _import_reference_values(cls):
        return {
            "source": set(Source.objects.filter(active=True).values_list("name", flat=True)),
            "product": set(Product.objects.filter(active=True).values_list("name", flat=True)),
            "learning_mode": set(LearningMode.objects.filter(active=True).values_list("name", flat=True)),
            "course_duration": set(CourseDuration.objects.filter(active=True).values_list("name", flat=True)),
            "status": set(Lead.Status.values),
        }

    @classmethod
    def _lookup_ci(cls, model, name):
        if not name:
            return None
        return model.objects.filter(name__iexact=name, active=True).first()

    @classmethod
    def _validate_import_row(cls, row, row_number, refs):
        errors = []
        name = str(row.get("name", "")).strip()
        email = str(row.get("email", "")).strip().lower()
        source_name = str(row.get("source", "")).strip()
        product_name = str(row.get("product", "")).strip()
        learning_mode_name = str(row.get("learning_mode", "")).strip()
        duration_name = str(row.get("course_duration", "")).strip()
        status_value = str(row.get("status", "")).strip().lower()

        # Name is the ONLY required import field.
        if not name:
            errors.append("Name is required")

        # Email is optional. Validate it only when provided.
        if email:
            try:
                validate_email(email)
            except DjangoValidationError:
                errors.append("Invalid email")

        # Source is optional. Validate it only when provided.
        if source_name and not any(
            source_name.lower() == value.lower()
            for value in refs["source"]
        ):
            errors.append("Unknown source")

        # All remaining master-data fields are optional.
        if product_name and not any(
            product_name.lower() == value.lower()
            for value in refs["product"]
        ):
            errors.append("Unknown course/product")

        if learning_mode_name and not any(
            learning_mode_name.lower() == value.lower()
            for value in refs["learning_mode"]
        ):
            errors.append("Unknown learning mode")

        if duration_name and not any(
            duration_name.lower() == value.lower()
            for value in refs["course_duration"]
        ):
            errors.append("Unknown course duration")

        if status_value and status_value not in {
            value.lower() for value in refs["status"]
        }:
            errors.append("Invalid status")

        return errors

    @classmethod
    def _find_import_duplicates(cls, email):
        if not email:
            return []
        duplicates = []
        leads = Lead.objects.filter(email__iexact=email, deleted_at__isnull=True).select_related("record__owner")
        for lead in leads:
            duplicates.append({"type": "lead", "id": lead.id, "name": lead.name, "email": lead.email})
        contacts = Contact.objects.filter(email__iexact=email, deleted_at__isnull=True).select_related("record__owner")
        for contact in contacts:
            duplicates.append({"type": "contact", "id": contact.id, "name": contact.name, "email": contact.email})
        return duplicates

    @classmethod
    def _normalise_duplicate_decisions(cls, request):
        raw = request.data.get("duplicate_decisions", {})
        if isinstance(raw, str):
            try:
                raw = json.loads(raw)
            except (TypeError, ValueError):
                raw = {}
        if not isinstance(raw, dict):
            raw = {}
        return {str(key): str(value).strip().lower() for key, value in raw.items()}

    @classmethod
    def _normalise_skip_rows(cls, request):
        raw = request.data.get("skip_rows", [])
        if isinstance(raw, str):
            try:
                raw = json.loads(raw)
            except (TypeError, ValueError):
                raw = [part.strip() for part in raw.split(",") if part.strip()]
        if not isinstance(raw, list):
            raw = []
        return {str(value) for value in raw}

    @action(detail=False, methods=["post"], url_path="import-preview")
    def import_preview(self, request):
        require_permission(request, "import_leads")
        upload = request.FILES.get("file")
        if not upload:
            return Response({"detail": "Upload a CSV or XLSX file."}, status=400)
        try:
            headers, data_rows = self._parse_import_upload(upload)
        except (ValueError, UnicodeDecodeError) as exc:
            return Response({"detail": str(exc)}, status=400)
        except Exception:
            return Response({"detail": "The file could not be read safely."}, status=400)

        mapping = self._suggest_import_mapping(headers)
        supplied_mapping = request.data.get("mapping")
        if supplied_mapping:
            if isinstance(supplied_mapping, str):
                try:
                    supplied_mapping = json.loads(supplied_mapping)
                except (TypeError, ValueError):
                    supplied_mapping = None
            if isinstance(supplied_mapping, dict):
                mapping = supplied_mapping

        refs = self._import_reference_values()
        preview = []
        all_errors = []
        duplicate_rows = []
        for row_number, raw_row in enumerate(data_rows, start=2):
            row = self._apply_import_mapping(raw_row, mapping)

            # Source is optional during import. When it is empty, preview the
            # active "Other" source that will be assigned during import.
            if not str(row.get("source", "")).strip():
                default_source = self._lookup_ci(Source, "Other")
                if default_source:
                    row["source"] = default_source.name

            errors = self._validate_import_row(row, row_number, refs)
            duplicates = self._find_import_duplicates(str(row.get("email", "")).strip().lower())
            if duplicates:
                duplicate_rows.append({"row": row_number, "duplicates": duplicates})
            item = {
                "row": row_number,
                "data": row,
                "errors": errors,
                "duplicates": duplicates,
            }
            if errors:
                all_errors.append({"row": row_number, "errors": errors})
            if len(preview) < 25:
                preview.append(item)

        accepted_values = {
            "source": sorted(refs["source"]),
            "product": sorted(refs["product"]),
            "learning_mode": sorted(refs["learning_mode"]),
            "course_duration": sorted(refs["course_duration"]),
            "status": sorted(refs["status"]),
        }
        return Response({
            "file_name": upload.name,
            "headers": headers,
            "mapping": mapping,
            "accepted_values": accepted_values,
            "preview": preview,
            "errors": all_errors,
            "duplicates": duplicate_rows,
            "total_rows": len(data_rows),
            "preview_rows": len(preview),
        })

    @action(detail=False, methods=["post"], url_path="import")
    def import_leads(self, request):
        require_permission(request, "import_leads")
        upload = request.FILES.get("file")
        if not upload:
            return Response({"detail": "Upload a CSV or XLSX file."}, status=400)

        try:
            headers, data_rows = self._parse_import_upload(upload)
        except (ValueError, UnicodeDecodeError) as exc:
            return Response({"detail": str(exc)}, status=400)
        except Exception:
            return Response({"detail": "The file could not be read safely."}, status=400)

        mapping = self._suggest_import_mapping(headers)
        supplied_mapping = request.data.get("mapping")
        if supplied_mapping:
            if isinstance(supplied_mapping, str):
                try:
                    supplied_mapping = json.loads(supplied_mapping)
                except (TypeError, ValueError):
                    supplied_mapping = None
            if isinstance(supplied_mapping, dict):
                mapping = supplied_mapping

        mode = str(request.data.get("mode", "preview")).strip().lower()
        if mode != "import":
            request._full_data = request.data.copy()
            return self.import_preview(request)

        refs = self._import_reference_values()
        duplicate_decisions = self._normalise_duplicate_decisions(request)
        skip_rows = self._normalise_skip_rows(request)
        prepared = []
        blocking_errors = []
        skipped_count = 0

        for row_number, raw_row in enumerate(data_rows, start=2):
            row = self._apply_import_mapping(raw_row, mapping)
            errors = self._validate_import_row(row, row_number, refs)
            email = str(row.get("email", "")).strip().lower()
            duplicates = self._find_import_duplicates(email)
            decision = duplicate_decisions.get(str(row_number), "")

            if duplicates and decision not in {"skip", "import_anyway", "update"}:
                errors = errors + ["Duplicate found; choose skip, import_anyway, or update"]
            if decision == "skip" or str(row_number) in skip_rows:
                skipped_count += 1
                prepared.append({"row": row_number, "data": row, "duplicates": duplicates, "decision": "skip", "errors": errors})
                continue
            if decision == "import_anyway" and duplicates:
                duplicates = []
            if decision == "update" and not duplicates:
                errors.append("No existing lead or contact was found to update")
            if errors:
                blocking_errors.append({"row": row_number, "errors": errors})
            prepared.append({"row": row_number, "data": row, "duplicates": duplicates, "decision": decision or "create", "errors": errors})

        if blocking_errors:
            return Response({
                "detail": "Nothing was imported. Fix the reported rows or explicitly skip them, then retry.",
                "errors": blocking_errors,
                "skipped": skipped_count,
                "total_rows": len(data_rows),
            }, status=400)

        batch_id = secrets.token_hex(8)
        created_count = 0
        updated_count = 0
        outcome_rows = []

        try:
            with transaction.atomic():
                batch = ImportBatch.objects.create(
                    batch_id=batch_id,
                    file_name=upload.name,
                    imported_by=request.user,
                    total_rows=len(data_rows),
                    skipped_count=skipped_count,
                )

                for item in prepared:
                    row_number = item["row"]
                    row = item["data"]
                    if item["decision"] == "skip":
                        outcome_rows.append({"row": row_number, "result": "skipped", "reason": item["errors"] or ["Skipped by user"]})
                        continue

                    email = str(row.get("email", "")).strip().lower()

                    # Source is optional in the uploaded file. If it is empty,
                    # automatically assign the active "Other" source so every
                    # imported Lead remains compatible with Contact conversion.
                    source = self._lookup_ci(Source, row.get("source"))
                    if not source:
                        source = self._lookup_ci(Source, "Other")
                    if not source:
                        raise ValidationError(
                            'Default Lead Source "Other" is not available or inactive.'
                        )

                    product = self._lookup_ci(Product, row.get("product"))
                    learning_mode = self._lookup_ci(LearningMode, row.get("learning_mode"))
                    course_duration = self._lookup_ci(CourseDuration, row.get("course_duration"))
                    status_value = str(row.get("status", "")).strip().lower() or Lead.Status.NOT_QUALIFIED
                    duplicate = item["duplicates"][0] if item["duplicates"] else None

                    if item["decision"] == "update" and duplicate:
                        if duplicate["type"] == "lead":
                            lead = Lead.objects.select_related("record").get(pk=duplicate["id"])
                            before = {
                                "name": lead.name,
                                "company": lead.company,
                                "email": lead.email,
                                "phone": lead.phone,
                                "source_id": lead.source_id,
                                "product_id": lead.product_id,
                                "learning_mode_id": lead.learning_mode_id,
                                "course_duration_id": lead.course_duration_id,
                                "town_city": lead.town_city,
                                "status": lead.status,
                                "notes": lead.notes,
                            }
                            lead.name = row.get("name", lead.name)
                            lead.company = row.get("company", lead.company)
                            lead.phone = row.get("phone", lead.phone)
                            lead.source = source
                            lead.product = product
                            lead.learning_mode = learning_mode
                            lead.course_duration = course_duration
                            lead.town_city = row.get("town_city", lead.town_city)
                            lead.status = status_value
                            lead.notes = row.get("notes", lead.notes)
                            lead.save(update_fields=["name", "company", "phone", "source", "product", "learning_mode", "course_duration", "town_city", "status", "notes"])
                            history(lead.record, request.user, "edited", {"changed": {"import_batch": {"from": None, "to": batch_id}, "row": {"from": None, "to": row_number}}})
                        else:
                            contact = Contact.objects.select_related("record").get(pk=duplicate["id"])
                            before = {
                                "name": contact.name,
                                "company": contact.company,
                                "email": contact.email,
                                "phone": contact.phone,
                                "source_id": contact.source_id,
                                "notes": contact.notes,
                            }
                            contact.name = row.get("name", contact.name)
                            contact.company = row.get("company", contact.company)
                            contact.phone = row.get("phone", contact.phone)
                            contact.source = source
                            contact.notes = row.get("notes", contact.notes)
                            contact.save(update_fields=["name", "company", "phone", "source", "notes"])
                            history(contact.record, request.user, "edited", {"changed": {"import_batch": {"from": None, "to": batch_id}, "row": {"from": None, "to": row_number}}})
                        updated_count += 1
                        outcome_rows.append({"row": row_number, "result": "updated", "record_type": duplicate["type"], "record_id": duplicate["id"]})
                        continue

                    lead = Lead.objects.create(
                        record=Record.objects.create(owner=request.user),
                        name=row.get("name", ""),
                        company=row.get("company", ""),
                        email=email,
                        phone=row.get("phone", ""),
                        source=source,
                        product=product,
                        learning_mode=learning_mode,
                        course_duration=course_duration,
                        town_city=row.get("town_city", ""),
                        status=status_value,
                        notes=row.get("notes", ""),
                    )
                    history(lead.record, request.user, "lead_captured", {"source": source.name if source else None, "import_batch": batch_id, "row": row_number})
                    created_count += 1
                    outcome_rows.append({"row": row_number, "result": "created", "record_type": "lead", "record_id": lead.id})

                batch.created_count = created_count
                batch.updated_count = updated_count
                batch.skipped_count = skipped_count
                batch.failed_count = 0
                batch.save(update_fields=["created_count", "updated_count", "skipped_count", "failed_count"])
                audit(request.user, "lead.import.batch", batch, after={
                    "batch_id": batch_id,
                    "file_name": upload.name,
                    "created": created_count,
                    "updated": updated_count,
                    "skipped": skipped_count,
                    "failed": 0,
                    "rows": outcome_rows,
                })
        except Exception as exc:
            return Response({"detail": f"The import was rolled back because the batch could not be completed: {exc}"}, status=400)

        return Response({
            "batch_id": batch_id,
            "file_name": upload.name,
            "created": created_count,
            "updated": updated_count,
            "skipped": skipped_count,
            "failed": 0,
            "total_rows": len(data_rows),
            "report_url": f"/api/leads/import-report/{batch_id}/",
        }, status=201)

    @action(detail=False, methods=["get"], url_path=r"import-report/(?P<batch_id>[0-9a-f]+)")
    def import_report(self, request, batch_id=None):
        require_permission(request, "import_leads")
        batch = ImportBatch.objects.filter(batch_id=batch_id).first()
        if not batch:
            return Response({"detail": "Import batch not found."}, status=404)
        if not request.user.has_crm_permission("view_workspace_records") and batch.imported_by_id != request.user.id:
            return Response(status=403)

        audit_row = batch.__class__
        from .models import AuditEvent
        event = AuditEvent.objects.filter(action="lead.import.batch", target_type="ImportBatch", target_id=str(batch.pk)).order_by("-created_at").first()
        rows = event.after.get("rows", []) if event else []

        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(["Batch ID", "File", "Row", "Result", "Record Type", "Record ID", "Reason"])
        for row in rows:
            writer.writerow([
                batch.batch_id,
                batch.file_name,
                row.get("row", ""),
                row.get("result", ""),
                row.get("record_type", ""),
                row.get("record_id", ""),
                "; ".join(row.get("reason", [])) if isinstance(row.get("reason"), list) else row.get("reason", ""),
            ])
        if not rows:
            writer.writerow([batch.batch_id, batch.file_name, "", "", "", "", "No row-level outcomes recorded"])
        writer.writerow([])
        writer.writerow(["Summary", "Created", batch.created_count, "Updated", batch.updated_count, "Skipped", batch.skipped_count, "Failed", batch.failed_count])

        response = HttpResponse(output.getvalue(), content_type="text/csv; charset=utf-8")
        response["Content-Disposition"] = f'attachment; filename="litecrm-import-{batch.batch_id}-report.csv"'
        return response

    @action(detail=False, methods=["post"], url_path=r"import-reverse/(?P<batch_id>[0-9a-f]+)")
    def import_reverse(self, request, batch_id=None):
        require_permission(request, "import_leads")
        batch = ImportBatch.objects.filter(batch_id=batch_id).first()
        if not batch:
            return Response({"detail": "Import batch not found."}, status=404)
        if batch.reversed_at:
            return Response({"detail": "This import batch has already been reversed."}, status=400)
        if not request.user.has_crm_permission("view_workspace_records") and batch.imported_by_id != request.user.id:
            return Response(status=403)

        events = list(batch.imported_by.audit_events.none()) if False else []
        from .models import HistoryEvent
        history_rows = list(HistoryEvent.objects.filter(detail__batch_id=batch_id).select_related("record", "record__lead", "record__contact").order_by("-created_at"))
        if not history_rows:
            return Response({"detail": "No reversible records were found for this batch."}, status=400)

        with transaction.atomic():
            reversed_count = 0
            for event in history_rows:
                detail = event.detail or {}
                operation = detail.get("operation")
                if operation == "create" and hasattr(event.record, "lead"):
                    lead = event.record.lead
                    if lead.deleted_at is None:
                        lead.deleted_at = timezone.now()
                        lead.save(update_fields=["deleted_at"])
                        reversed_count += 1
                elif operation == "create" and hasattr(event.record, "contact"):
                    contact = event.record.contact
                    if contact.deleted_at is None:
                        contact.deleted_at = timezone.now()
                        contact.save(update_fields=["deleted_at"])
                        reversed_count += 1
                elif operation == "update":
                    before = detail.get("before", {})
                    if hasattr(event.record, "lead"):
                        lead = event.record.lead
                        for field, value in before.items():
                            setattr(lead, field, value)
                        lead.save(update_fields=list(before.keys()))
                        reversed_count += 1
                    elif hasattr(event.record, "contact"):
                        contact = event.record.contact
                        for field, value in before.items():
                            setattr(contact, field, value)
                        contact.save(update_fields=list(before.keys()))
                        reversed_count += 1
            batch.reversed_at = timezone.now()
            batch.save(update_fields=["reversed_at"])
            audit(request.user, "lead.import.batch_reversed", batch, after={"batch_id": batch_id, "reversed_records": reversed_count})

        return Response({"batch_id": batch_id, "reversed": reversed_count, "reversed_at": batch.reversed_at})


class RoleViewSet(viewsets.ViewSet):
    """Admin-only API for viewing and managing CRM roles and permissions."""

    ROLE_KEYS = {"admin", "sales_manager", "salesperson"}

    def _require_permission(self, request):
        require_permission(request, "manage_permissions")

    def list(self, request):
        self._require_permission(request)
        roles = Role.objects.all().order_by("id")
        return Response(RoleSerializer(roles, many=True).data)

    def retrieve(self, request, pk=None):
        self._require_permission(request)
        role = Role.objects.filter(pk=pk).first()
        if not role:
            return Response({"detail": "Role not found."}, status=404)
        return Response(RoleSerializer(role).data)

    def create(self, request):
        self._require_permission(request)
        key = str(request.data.get("key", "")).strip().lower()
        name = str(request.data.get("name", "")).strip()
        permissions = request.data.get("permissions", [])

        if not key:
            return Response({"key": "Role key is required."}, status=400)
        if key not in self.ROLE_KEYS:
            return Response({"key": "Choose admin, sales_manager, or salesperson."}, status=400)
        if not name:
            return Response({"name": "Role name is required."}, status=400)
        if not isinstance(permissions, list):
            return Response({"permissions": "Permissions must be a list."}, status=400)

        valid_permissions = {value for value, _ in Permission.choices}
        invalid = sorted(set(permissions) - valid_permissions)
        if invalid:
            return Response({"permissions": f"Invalid permissions: {', '.join(invalid)}."}, status=400)
        if Role.objects.filter(key=key).exists():
            return Response({"key": "A role with this key already exists."}, status=400)

        role = Role.objects.create(key=key, name=name, permissions=permissions)
        audit(request.user, "role.created", role, after=RoleSerializer(role).data)
        return Response(RoleSerializer(role).data, status=201)

    def update(self, request, pk=None):
        return self._update(request, pk, partial=False)

    def partial_update(self, request, pk=None):
        return self._update(request, pk, partial=True)

    def _update(self, request, pk=None, partial=False):
        self._require_permission(request)
        role = Role.objects.filter(pk=pk).first()
        if not role:
            return Response({"detail": "Role not found."}, status=404)

        data = request.data
        if "key" in data:
            key = str(data.get("key", "")).strip().lower()
            if key not in self.ROLE_KEYS:
                return Response({"key": "Choose admin, sales_manager, or salesperson."}, status=400)
            if Role.objects.filter(key=key).exclude(pk=role.pk).exists():
                return Response({"key": "A role with this key already exists."}, status=400)
            role.key = key
        elif not partial:
            return Response({"key": "Role key is required."}, status=400)

        if "name" in data:
            name = str(data.get("name", "")).strip()
            if not name:
                return Response({"name": "Role name is required."}, status=400)
            role.name = name
        elif not partial:
            return Response({"name": "Role name is required."}, status=400)

        if "permissions" in data:
            permissions = data.get("permissions")
            if not isinstance(permissions, list):
                return Response({"permissions": "Permissions must be a list."}, status=400)
            valid_permissions = {value for value, _ in Permission.choices}
            invalid = sorted(set(permissions) - valid_permissions)
            if invalid:
                return Response({"permissions": f"Invalid permissions: {', '.join(invalid)}."}, status=400)
            role.permissions = permissions

        role.save()
        audit(request.user, "role.updated", role, after=RoleSerializer(role).data)
        return Response(RoleSerializer(role).data)


class UserViewSet(viewsets.ViewSet):
    """Admin user-management API plus active-user lookup for CRM assignments."""

    ROLE_KEYS = {"admin", "sales_manager", "salesperson"}

    def _require_admin(self, request):
        if not request.user.has_crm_permission("manage_users"):
            raise PermissionDenied("Only Admin users can manage users.")

    @staticmethod
    def _serialize_user(user):
        return {
            "id": user.id,
            "username": user.username,
            "display_name": user.get_full_name() or user.username or user.email,
            "first_name": user.first_name,
            "last_name": user.last_name,
            "email": user.email,
            "mobile": user.mobile,
            "role": user.role.key if user.role else None,
            "role_name": user.role.name if user.role else None,
            "is_active": user.is_active,
            "must_change_password": user.must_change_password,
        }

    def list(self, request):
        # Admins need the complete user list for User Management.
        # Sales Managers / Salespersons only need the active-user subset when
        # they have permission to reassign leads. This keeps the user-management
        # API protected while still allowing the Lead Owner dropdown to hydrate.
        if request.user.has_crm_permission("manage_users"):
            users = User.objects.select_related("role").order_by(
                "first_name", "last_name", "email"
            )
        elif request.user.has_crm_permission("reassign_lead"):
            users = User.objects.select_related("role").filter(
                is_active=True
            ).order_by("first_name", "last_name", "email")
        else:
            users = User.objects.select_related("role").filter(
        pk=request.user.pk,
        is_active=True,
    )

        return Response([self._serialize_user(user) for user in users])

    def retrieve(self, request, pk=None):
        self._require_admin(request)
        user = User.objects.select_related("role").filter(pk=pk).first()
        if not user:
            return Response({"detail": "User not found."}, status=404)
        return Response(self._serialize_user(user))

    def create(self, request):
        self._require_admin(request)
        data = request.data
        email = str(data.get("email", "")).strip().lower()
        username = str(data.get("username") or email).strip()
        first_name = str(data.get("first_name", "")).strip()
        last_name = str(data.get("last_name", "")).strip()
        mobile = str(data.get("mobile") or "").strip()
        role_key = str(data.get("role", "")).strip()
        password = str(data.get("password", ""))

        if not email:
            return Response({"email": "Email is required."}, status=400)
        try:
            validate_email(email)
        except DjangoValidationError:
            return Response({"email": "Enter a valid email address."}, status=400)
        if not first_name:
            return Response({"first_name": "First name is required."}, status=400)
        if not role_key:
            return Response({"role": "Role is required."}, status=400)
        if role_key not in self.ROLE_KEYS:
            return Response({"role": "Choose Admin, Sales Manager, or Salesperson."}, status=400)
        if not password:
            return Response({"password": "Password is required."}, status=400)
        if User.objects.filter(username__iexact=username).exists():
            return Response({"username": "A user with this username already exists."}, status=400)
        if User.objects.filter(email__iexact=email).exists():
            return Response({"email": "A user with this email already exists."}, status=400)
        if mobile and User.objects.filter(mobile=mobile).exists():
            return Response({"mobile": "A user with this mobile number already exists."}, status=400)

        try:
            validate_password(password)
        except Exception as exc:
            return Response({"password": "; ".join(exc.messages)}, status=400)
        if not any(c.isupper() for c in password) or not any(not c.isalnum() for c in password):
            return Response({"password": "Password must contain an uppercase letter and a special character."}, status=400)

        role = Role.objects.filter(key=role_key).first()
        if not role:
            return Response({"role": "The selected role has not been configured."}, status=400)

        user = User(
            username=username,
            email=email,
            first_name=first_name,
            last_name=last_name,
            mobile=mobile or None,
            role=role,
            is_active=bool(data.get("is_active", True)),
            must_change_password=True,
        )
        user.set_password(password)
        user.save()

        audit(request.user, "user.created", user, after={"username": user.username, "email": user.email, "role": role.key, "is_active": user.is_active})
        return Response(self._serialize_user(user), status=201)

    def partial_update(self, request, pk=None):
        return self._update_user(request, pk, partial=True)

    def update(self, request, pk=None):
        return self._update_user(request, pk, partial=False)

    def _update_user(self, request, pk, partial=False):
        self._require_admin(request)
        user = User.objects.select_related("role").filter(pk=pk).first()
        if not user:
            return Response({"detail": "User not found."}, status=404)

        data = request.data
        if not partial or "first_name" in data:
            user.first_name = str(data.get("first_name", "")).strip()
        if not partial or "last_name" in data:
            user.last_name = str(data.get("last_name", "")).strip()
        if "email" in data:
            email = str(data.get("email", "")).strip().lower()
            if not email:
                return Response({"email": "Email is required."}, status=400)
            try:
                validate_email(email)
            except DjangoValidationError:
                return Response({"email": "Enter a valid email address."}, status=400)
            if User.objects.filter(email__iexact=email).exclude(pk=user.pk).exists():
                return Response({"email": "A user with this email already exists."}, status=400)
            user.email = email
        if "username" in data:
            username = str(data.get("username", "")).strip()
            if not username:
                return Response({"username": "Username is required."}, status=400)
            if User.objects.filter(username__iexact=username).exclude(pk=user.pk).exists():
                return Response({"username": "A user with this username already exists."}, status=400)
            user.username = username
        if "mobile" in data:
            mobile = str(data.get("mobile") or "").strip()
            if mobile and User.objects.filter(mobile=mobile).exclude(pk=user.pk).exists():
                return Response({"mobile": "A user with this mobile number already exists."}, status=400)
            user.mobile = mobile or None
        if "role" in data:
            role_key = str(data.get("role", "")).strip()
            if role_key not in self.ROLE_KEYS:
                return Response({"role": "Choose Admin, Sales Manager, or Salesperson."}, status=400)
            if user.id == request.user.id and role_key != (user.role.key if user.role else None):
                return Response({"role": "You cannot change your own role."}, status=400)
            role = Role.objects.filter(key=role_key).first()
            if not role:
                return Response({"role": "The selected role has not been configured."}, status=400)
            user.role = role
        if "is_active" in data:
            new_active = bool(data.get("is_active"))
            if user.id == request.user.id and not new_active:
                return Response({"is_active": "You cannot deactivate your own account."}, status=400)
            user.is_active = new_active
        user.save()

        if not user.is_active:
            _revoke_sessions(user)

        audit(request.user, "user.updated", user)
        return Response(self._serialize_user(user))

    @action(detail=True, methods=["post"], url_path="password")
    def change_password(self, request, pk=None):
        self._require_admin(request)
        user = User.objects.filter(pk=pk).first()
        if not user:
            return Response({"detail": "User not found."}, status=404)
        if user.id == request.user.id:
            return Response({"detail": "Use your own Change Password option."}, status=400)

        new_password = str(request.data.get("new_password", ""))
        if not new_password:
            return Response({"new_password": "New password is required."}, status=400)
        try:
            validate_password(new_password, user)
        except Exception as exc:
            return Response({"new_password": "; ".join(exc.messages)}, status=400)
        if not any(c.isupper() for c in new_password) or not any(not c.isalnum() for c in new_password):
            return Response({"new_password": "Password must contain an uppercase letter and a special character."}, status=400)

        user.set_password(new_password)
        user.must_change_password = True
        user.save(update_fields=["password", "must_change_password"])
        _revoke_sessions(user)
        audit(request.user, "auth.password.admin_reset", user)
        return Response({"detail": "Password changed. The user must change it at next sign-in."})

    @action(detail=True, methods=["patch"], url_path="status")
    def change_status(self, request, pk=None):
        self._require_admin(request)
        user = User.objects.filter(pk=pk).first()
        if not user:
            return Response({"detail": "User not found."}, status=404)
        if "is_active" not in request.data:
            return Response({"is_active": "is_active is required."}, status=400)
        active = bool(request.data.get("is_active"))
        if user.id == request.user.id and not active:
            return Response({"is_active": "You cannot deactivate your own account."}, status=400)
        user.is_active = active
        user.save(update_fields=["is_active"])
        if not active:
            _revoke_sessions(user)
        audit(request.user, "user.status_changed", user, after={"is_active": active})
        return Response(self._serialize_user(user))

    @action(detail=True, methods=["patch"], url_path="role")
    def change_role(self, request, pk=None):
        self._require_admin(request)
        user = User.objects.select_related("role").filter(pk=pk).first()
        if not user:
            return Response({"detail": "User not found."}, status=404)
        if user.id == request.user.id:
            return Response({"detail": "You cannot change your own role."}, status=400)
        role_key = str(request.data.get("role", "")).strip()
        if role_key not in self.ROLE_KEYS:
            return Response({"role": "Choose Admin, Sales Manager, or Salesperson."}, status=400)
        role = Role.objects.filter(key=role_key).first()
        if not role:
            return Response({"role": "The selected role has not been configured."}, status=400)
        old_role = user.role.key if user.role else None
        user.role = role
        user.save(update_fields=["role"])
        audit(request.user, "user.role_changed", user, before={"role": old_role}, after={"role": role.key})
        return Response(self._serialize_user(user))


class ContactViewSet(viewsets.ModelViewSet):
    serializer_class = ContactSerializer

    def retrieve(self, request, *args, **kwargs):
        require_permission(request, "view_contacts")
        return super().retrieve(request, *args, **kwargs)

    def update(self, request, *args, **kwargs):
        require_permission(request, "edit_contact")
        return self._update_contact(request, partial=False, *args, **kwargs)

    def partial_update(self, request, *args, **kwargs):
        require_permission(request, "edit_contact")
        return self._update_contact(request, partial=True, *args, **kwargs)

    def _update_contact(self, request, *args, partial=False, **kwargs):
        contact = self.get_object()
        payload = request.data.copy()
        owner_value = payload.pop("owner", None)
        new_owner = None
        if owner_value is not None:
            try:
                owner_id = int(owner_value)
            except (TypeError, ValueError):
                return Response({"owner": "Invalid owner."}, status=400)
            if owner_id != contact.record.owner_id and not request.user.has_crm_permission("reassign_contact"):
                return Response({"detail": "You do not have permission to reassign this contact."}, status=403)
            new_owner = User.objects.filter(pk=owner_id, is_active=True).first()
            if not new_owner:
                return Response({"owner": "Owner must be an active user."}, status=400)
        before = self.get_serializer(contact).data
        serializer = self.get_serializer(contact, data=payload, partial=partial)
        serializer.is_valid(raise_exception=True)
        updated = serializer.save()
        if new_owner:
            old_owner = contact.record.owner_id
            updated.record.owner = new_owner
            updated.record.save(update_fields=["owner"])
            history(
                updated.record,
                request.user,
                "reassigned",
                {
                    "from_owner": (User.objects.filter(pk=old_owner).first().get_full_name() if User.objects.filter(pk=old_owner).first() else str(old_owner)),
                    "to_owner": new_owner.get_full_name() or new_owner.username,
                },
            )
        after = self.get_serializer(updated).data
        changed = _history_changes(
            before,
            after,
            ("name", "company", "email", "phone", "source", "product", "learning_mode", "course_duration", "town_city", "status", "disqualification_reason", "notes", "industry", "location"),
        )
        if changed:
            history(updated.record, request.user, "edited", {"changed": changed})
        audit(request.user, "contact.updated", updated, before=before, after=after)
        return Response(after)

    @transaction.atomic
    def destroy(self, request, *args, **kwargs):
        require_permission(request, "delete_contact")
        contact = self.get_object()
        record = contact.record
        response = super().destroy(request, *args, **kwargs)
        if response.status_code == status.HTTP_204_NO_CONTENT:
            history(record, request.user, "deleted", {"entity": "Contact", "name": contact.name, "company": contact.company})
        return response

    def get_queryset(self):
        return Contact.objects.filter(record__in=visible_records(self.request.user), deleted_at__isnull=True).select_related("record__owner", "source", "product", "learning_mode", "course_duration")

    def list(self, request, *args, **kwargs):
        if not request.user.has_crm_permission("view_contacts"):
            return Response(status=403)
        return super().list(request, *args, **kwargs)

    def create(self, request, *args, **kwargs):
        if not request.user.has_crm_permission("create_contact"):
            return Response(status=403)
        payload = request.data.copy()
        owner_value = payload.pop("owner", None) or request.user.id
        try:
            owner_id = int(owner_value)
        except (TypeError, ValueError):
            return Response({"owner": "Invalid owner."}, status=400)
        if owner_id != request.user.id and not request.user.has_crm_permission("reassign_contact"):
            return Response({"detail": "You do not have permission to assign this contact to another user."}, status=403)
        owner = User.objects.filter(pk=owner_id, is_active=True).first()
        if not owner:
            return Response({"owner": "Owner must be an active user."}, status=400)
        serializer = self.get_serializer(data=payload)
        serializer.is_valid(raise_exception=True)
        contact = serializer.save(record=Record.objects.create(owner=owner))
        audit(request.user, "contact.created", contact, after=serializer.data)
        return Response(self.get_serializer(contact).data, status=201)
    @action(detail=True, methods=["get", "post"])
    def activities(self, request, pk=None):
        contact = self.get_object(); assert_visible(request.user, contact.record)
        if request.method == "GET":
            require_permission(request, "view_activity")
            activities = contact.record.activities.select_related("author", "deal").order_by("-happened_at", "-id")

            try:
                page_size = int(request.query_params.get("page_size", 20))
            except (TypeError, ValueError):
                page_size = 20
            page_size = min(max(page_size, 1), 50)

            try:
                page_number = int(request.query_params.get("page", 1))
            except (TypeError, ValueError):
                page_number = 1
            page_number = max(page_number, 1)

            paginator = Paginator(activities, page_size)
            page = paginator.get_page(page_number)

            return Response({
                "count": paginator.count,
                "page": page.number,
                "page_size": page_size,
                "next": page.next_page_number() if page.has_next() else None,
                "previous": page.previous_page_number() if page.has_previous() else None,
                "results": ActivitySerializer(page.object_list, many=True, context={"request": request}).data,
            })
        if not request.user.has_crm_permission("log_activity"): return Response(status=403)
        serializer = ActivitySerializer(data=request.data); serializer.is_valid(raise_exception=True)
        activity = serializer.save(record=contact.record, author=request.user)
        if not contact.record.last_contact_at or activity.happened_at > contact.record.last_contact_at:
            contact.record.last_contact_at = activity.happened_at; contact.record.save(update_fields=["last_contact_at"])
        return Response(ActivitySerializer(activity, context={"request": request}).data, status=201)
    @action(detail=True, methods=["get"])
    def history(self, request, pk=None):
        contact = self.get_object()
        if not can_read_record_history(request.user, contact.record):
            return Response(status=status.HTTP_403_FORBIDDEN)
        return Response(HistorySerializer(contact.record.history.all(), many=True).data)
    @action(detail=True, methods=["post"])
    def deals(self, request, pk=None):
        if not request.user.has_crm_permission("create_deal"): return Response(status=403)
        contact = self.get_object(); assert_visible(request.user, contact.record)
        # Enforce the date rule at the contact-to-deal endpoint itself, before
        # serializer processing, so a past date cannot slip through this path.
        raw_close_date = request.data.get("expected_close_date")
        if raw_close_date:
            try:
                submitted_close_date = date.fromisoformat(str(raw_close_date)[:10])
            except (TypeError, ValueError):
                return Response(
                    {"expected_close_date": "Enter a valid date."},
                    status=400,
                )
            if submitted_close_date < timezone.localdate():
                return Response(
                    {"expected_close_date": "Expected close date cannot be in the past."},
                    status=400,
                )
        serializer = DealSerializer(data=request.data); serializer.is_valid(raise_exception=True)
        owner_id = serializer.validated_data.pop("owner_id", None)
        record = contact.record
        if owner_id is not None and owner_id != record.owner_id:
            record.owner_id = owner_id
        record.last_contact_at = timezone.now()
        record.save(update_fields=["owner", "last_contact_at"] if owner_id is not None else ["last_contact_at"])
        deal = serializer.save(record=record, stage=Deal.Stage.PROPOSAL)
        history(record, request.user, "deal_created", {"product": deal.product.name, "value_minor": deal.value_minor, "stage": deal.stage, "expected_close_date": deal.expected_close_date.isoformat()})
        return Response(DealSerializer(deal).data, status=201)


class DealViewSet(viewsets.ModelViewSet):
    serializer_class = DealSerializer

    def create(self, request, *args, **kwargs):
        return Response({"detail": "Deals can only be created from a contact."}, status=405)

    def retrieve(self, request, *args, **kwargs):
        require_permission(request, "view_deals")
        return super().retrieve(request, *args, **kwargs)

    def update(self, request, *args, **kwargs):
        require_permission(request, "edit_deal")
        return self._update_deal(request, *args, partial=False, **kwargs)

    def partial_update(self, request, *args, **kwargs):
        require_permission(request, "edit_deal")
        return self._update_deal(request, *args, partial=True, **kwargs)

    def _update_deal(self, request, *args, partial=False, **kwargs):
        deal = self.get_object()
        before = self.get_serializer(deal).data
        # Validate close dates explicitly on updates as well as in the serializer.
        if "expected_close_date" in request.data and request.data.get("expected_close_date"):
            try:
                submitted_close_date = date.fromisoformat(str(request.data.get("expected_close_date"))[:10])
            except (TypeError, ValueError):
                return Response(
                    {"expected_close_date": "Enter a valid date."},
                    status=400,
                )
            if submitted_close_date < timezone.localdate():
                return Response(
                    {"expected_close_date": "Expected close date cannot be in the past."},
                    status=400,
                )
        if "value_minor" in request.data:
            if deal.stage in (Deal.Stage.WON, Deal.Stage.LOST):
                return Response({"detail": "The value is locked because this deal is closed."}, status=400)
            if not request.user.has_crm_permission("edit_deal_value"):
                return Response(status=403)
        serializer = self.get_serializer(deal, data=request.data, partial=partial)
        serializer.is_valid(raise_exception=True)
        updated = serializer.save()
        after = self.get_serializer(updated).data
        if before.get("value_minor") != after.get("value_minor"):
            history(
                updated.record,
                request.user,
                "deal_value_updated",
                {
                    "from_value_minor": before.get("value_minor"),
                    "to_value_minor": after.get("value_minor"),
                },
            )
        changed = _history_changes(
            before,
            after,
            ("product", "expected_close_date", "loss_reason", "notes"),
        )
        if changed:
            history(updated.record, request.user, "edited", {"changed": changed})
        return Response(after)

    def destroy(self, request, *args, **kwargs):
        require_permission(request, "delete_deal")
        deal = self.get_object()
        record = deal.record
        response = super().destroy(request, *args, **kwargs)
        if response.status_code == status.HTTP_204_NO_CONTENT:
            history(record, request.user, "deleted", {"entity": "Deal", "product": deal.product.name})
        return response

    def get_queryset(self):
        return (
            Deal.objects
            .filter(record__in=visible_records(self.request.user))
            .select_related("record__owner", "product")
            .order_by("-created_at", "-id")
        )

    def list(self, request, *args, **kwargs):
        if not request.user.has_crm_permission("view_deals"):
            return Response(status=403)
        return super().list(request, *args, **kwargs)

    @action(detail=True, methods=["post"])
    def move(self, request, pk=None):
        deal = self.get_object()
        stage = request.data.get("stage")
        closed = {Deal.Stage.WON, Deal.Stage.LOST}
        needed = "close_won" if stage == Deal.Stage.WON else "close_lost" if stage == Deal.Stage.LOST else "move_deal"
        if not request.user.has_crm_permission(needed):
            return Response(status=403)
        if stage not in Deal.Stage.values:
            return Response({"stage": "Invalid pipeline stage."}, status=400)
        if stage == Deal.Stage.LOST and not request.data.get("loss_reason"):
            return Response({"loss_reason": "A loss reason is required."}, status=400)
        old = deal.stage
        deal.stage = stage
        deal.loss_reason = request.data.get("loss_reason", deal.loss_reason)
        deal.closed_at = timezone.now() if stage in closed else None
        deal.save()
        if old != stage:
            detail = {
                "from_stage": old,
                "to_stage": stage,
                "value_minor": deal.value_minor,
                "loss_reason": deal.loss_reason if stage == Deal.Stage.LOST else None,
            }
            history(deal.record, request.user, "deal_stage_changed", detail)
            if stage == Deal.Stage.WON:
                history(deal.record, request.user, "deal_won", detail)
            elif stage == Deal.Stage.LOST:
                history(deal.record, request.user, "deal_lost", detail)
        return Response(DealSerializer(deal).data)


class ProductViewSet(viewsets.ModelViewSet):
    queryset = Product.objects.all().order_by("name")
    serializer_class = ProductSerializer

    def _can_manage_catalogue(self):
        return self.request.user.has_crm_permission("manage_catalogue")

    def get_queryset(self):
        # Admin/catalogue managers can see all products. Lead users only need
        # active products so the Course Interested dropdown works for them.
        if self._can_manage_catalogue():
            return Product.objects.all().order_by("name")

        if self.request.user.has_crm_permission("view_leads"):
            return Product.objects.filter(active=True).order_by("name")

        raise PermissionDenied(
            "You do not have permission to view the product catalogue."
        )

    def list(self, request, *args, **kwargs):
        return super().list(request, *args, **kwargs)

    def retrieve(self, request, *args, **kwargs):
        return super().retrieve(request, *args, **kwargs)

    def perform_create(self, serializer):
        if not self._can_manage_catalogue():
            raise PermissionDenied("You do not have permission to manage the product catalogue.")
        serializer.save()

    def perform_update(self, serializer):
        if not self._can_manage_catalogue():
            raise PermissionDenied("You do not have permission to manage the product catalogue.")
        serializer.save()

    def destroy(self, request, *args, **kwargs):
        if not self._can_manage_catalogue():
            raise PermissionDenied("You do not have permission to manage the product catalogue.")
        instance = self.get_object()
        try:
            instance.delete()
        except ProtectedError:
            return Response(
                {"detail": "This course is used by existing Leads or Contacts and cannot be deleted. Deactivate it instead to preserve existing records."},
                status=status.HTTP_409_CONFLICT,
            )
        return Response(status=status.HTTP_204_NO_CONTENT)


class LearningModeViewSet(viewsets.ModelViewSet):
    serializer_class = LearningModeSerializer

    def get_queryset(self):
        if self.request.user.has_crm_permission("manage_catalogue"):
            return LearningMode.objects.all().order_by("name")
        if not self.request.user.has_crm_permission("view_leads"):
            raise PermissionDenied("You do not have permission to view learning modes.")
        return LearningMode.objects.filter(active=True).order_by("name")

    def perform_create(self, serializer):
        if not self.request.user.has_crm_permission("manage_catalogue"):
            raise PermissionDenied("You do not have permission to manage learning modes.")
        serializer.save()

    def perform_update(self, serializer):
        if not self.request.user.has_crm_permission("manage_catalogue"):
            raise PermissionDenied("You do not have permission to manage learning modes.")
        serializer.save()

    def destroy(self, request, *args, **kwargs):
        if not request.user.has_crm_permission("manage_catalogue"):
            raise PermissionDenied("You do not have permission to manage learning modes.")
        instance = self.get_object()
        try:
            instance.delete()
        except ProtectedError:
            return Response(
                {"detail": "This Learning Mode is used by existing Leads or Contacts and cannot be deleted. Deactivate it instead to preserve existing records."},
                status=status.HTTP_409_CONFLICT,
            )
        return Response(status=status.HTTP_204_NO_CONTENT)


class CourseDurationViewSet(viewsets.ModelViewSet):
    serializer_class = CourseDurationSerializer

    def get_queryset(self):
        if self.request.user.has_crm_permission("manage_catalogue"):
            return CourseDuration.objects.all().order_by("name")
        if not self.request.user.has_crm_permission("view_leads"):
            raise PermissionDenied("You do not have permission to view course durations.")
        return CourseDuration.objects.filter(active=True).order_by("name")

    def perform_create(self, serializer):
        if not self.request.user.has_crm_permission("manage_catalogue"):
            raise PermissionDenied("You do not have permission to manage course durations.")
        serializer.save()

    def perform_update(self, serializer):
        if not self.request.user.has_crm_permission("manage_catalogue"):
            raise PermissionDenied("You do not have permission to manage course durations.")
        serializer.save()

    def destroy(self, request, *args, **kwargs):
        if not request.user.has_crm_permission("manage_catalogue"):
            raise PermissionDenied("You do not have permission to manage course durations.")
        instance = self.get_object()
        try:
            instance.delete()
        except ProtectedError:
            return Response(
                {"detail": "This Course Duration is used by existing Leads or Contacts and cannot be deleted. Deactivate it instead to preserve existing records."},
                status=status.HTTP_409_CONFLICT,
            )
        return Response(status=status.HTTP_204_NO_CONTENT)


class SourceViewSet(viewsets.ModelViewSet):
    serializer_class = SourceSerializer

    def get_queryset(self):
        return Source.objects.filter(
            active=True
        ).order_by("name")

    def _can_manage_catalogue(self):
        return self.request.user.has_crm_permission(
            "manage_catalogue"
        )

    def create(self, request, *args, **kwargs):
        if not self._can_manage_catalogue():
            return Response(status=403)
        return super().create(request, *args, **kwargs)

    def update(self, request, *args, **kwargs):
        if not self._can_manage_catalogue():
            return Response(status=403)
        return super().update(request, *args, **kwargs)

    def partial_update(self, request, *args, **kwargs):
        if not self._can_manage_catalogue():
            return Response(status=403)
        return super().partial_update(request, *args, **kwargs)

    def destroy(self, request, *args, **kwargs):
        if not self._can_manage_catalogue():
            return Response(status=403)

        instance = self.get_object()

        try:
            instance.delete()
        except ProtectedError:
            return Response(
                {
                    "detail": (
                        "This lead source is already used by existing Leads "
                        "and cannot be deleted."
                    )
                },
                status=status.HTTP_409_CONFLICT,
            )

        return Response(status=status.HTTP_204_NO_CONTENT)


class NotificationViewSet(viewsets.ViewSet):
    """Authenticated user's own in-app notifications."""

    def list(self, request):
        notifications = Notification.objects.filter(
            user=request.user
        ).order_by("-created_at", "-id")[:50]
        return Response(NotificationSerializer(notifications, many=True).data)

    @action(detail=False, methods=["post"], url_path="mark-all-read")
    def mark_all_read(self, request):
        updated = Notification.objects.filter(
            user=request.user,
            read_at__isnull=True,
        ).update(read_at=timezone.now())
        return Response({"updated": updated})


class RecordViewSet(viewsets.ViewSet):
    """Read/write the one shared timeline for a lead, contact, or deal."""
    def get_record(self, request, pk):
        record = Record.objects.get(pk=pk); assert_visible(request.user, record); return record
    @action(detail=True, methods=["get", "post"])
    def activities(self, request, pk=None):
        record = self.get_record(request, pk)
        if request.method == "GET":
            require_permission(request, "view_activity")
            activities = record.activities.select_related("author", "deal").order_by("-happened_at", "-id")

            try:
                page_size = int(request.query_params.get("page_size", 20))
            except (TypeError, ValueError):
                page_size = 20
            page_size = min(max(page_size, 1), 50)

            try:
                page_number = int(request.query_params.get("page", 1))
            except (TypeError, ValueError):
                page_number = 1
            page_number = max(page_number, 1)

            paginator = Paginator(activities, page_size)
            page = paginator.get_page(page_number)

            return Response({
                "count": paginator.count,
                "page": page.number,
                "page_size": page_size,
                "next": page.next_page_number() if page.has_next() else None,
                "previous": page.previous_page_number() if page.has_previous() else None,
                "results": ActivitySerializer(page.object_list, many=True, context={"request": request}).data,
            })
        if not request.user.has_crm_permission("log_activity"): return Response(status=403)
        serializer = ActivitySerializer(data=request.data); serializer.is_valid(raise_exception=True)
        deal_id = request.data.get("deal")
        if deal_id and not record.deals.filter(id=deal_id).exists(): return Response({"deal": "This deal does not belong to the record."}, status=400)
        activity = serializer.save(record=record, author=request.user)
        if not record.last_contact_at or activity.happened_at > record.last_contact_at:
            record.last_contact_at = activity.happened_at; record.save(update_fields=["last_contact_at"])
        audit(request.user, "activity.created", activity, after={"type": activity.type, "subject": activity.subject})
        return Response(ActivitySerializer(activity, context={"request": request}).data, status=201)
    def _get_activity_for_record(self, request, record, activity_id):
        try:
            activity = record.activities.select_related("author", "deal", "voided_by").get(pk=activity_id)
        except Activity.DoesNotExist:
            return None
        return activity

    @action(
        detail=True,
        methods=["patch"],
        url_path=r"activities/(?P<activity_id>[^/.]+)/edit",
    )
    def edit_activity(self, request, pk=None, activity_id=None):
        record = self.get_record(request, pk)
        require_permission(request, "correct_activity")
        activity = self._get_activity_for_record(request, record, activity_id)
        if activity is None:
            return Response({"detail": "Activity not found."}, status=status.HTTP_404_NOT_FOUND)
        if activity.voided_at:
            return Response({"detail": "A voided activity cannot be edited."}, status=status.HTTP_400_BAD_REQUEST)
        if activity.author_id != request.user.id:
            return Response({"detail": "Only the activity author can edit this activity."}, status=status.HTTP_403_FORBIDDEN)
        window = timedelta(hours=max(settings.ACTIVITY_CORRECTION_WINDOW_HOURS, 0))
        if timezone.now() > activity.created_at + window:
            return Response({"detail": "The correction window for this activity has expired."}, status=status.HTTP_400_BAD_REQUEST)

        subject = request.data.get("subject", activity.subject)
        description = request.data.get("description", activity.description)
        if not isinstance(subject, str) or not subject.strip():
            return Response({"subject": "Subject is required."}, status=status.HTTP_400_BAD_REQUEST)

        if activity.original_subject is None:
            activity.original_subject = activity.subject
        if activity.original_description is None:
            activity.original_description = activity.description
        activity.subject = subject.strip()
        activity.description = description if isinstance(description, str) else activity.description
        activity.edited_at = timezone.now()
        activity.save(update_fields=["original_subject", "original_description", "subject", "description", "edited_at"])
        history(record, request.user, "activity_edited", {
            "activity_id": activity.id,
            "original_subject": activity.original_subject,
            "subject": activity.subject,
            "edited_at": activity.edited_at.isoformat(),
        })
        audit(request.user, "activity.edited", activity, after={"subject": activity.subject, "description": activity.description})
        return Response(ActivitySerializer(activity, context={"request": request}).data)

    @action(
        detail=True,
        methods=["post"],
        url_path=r"activities/(?P<activity_id>[^/.]+)/void",
    )
    def void_activity(self, request, pk=None, activity_id=None):
        record = self.get_record(request, pk)
        require_permission(request, "correct_activity")
        activity = self._get_activity_for_record(request, record, activity_id)
        if activity is None:
            return Response({"detail": "Activity not found."}, status=status.HTTP_404_NOT_FOUND)
        if activity.voided_at:
            return Response({"detail": "Activity is already voided."}, status=status.HTTP_400_BAD_REQUEST)
        is_admin = request.user.is_superuser or getattr(getattr(request.user, "role", None), "key", None) == "admin"
        if activity.author_id != request.user.id and not is_admin:
            return Response({"detail": "Only the activity author or an Admin can void this activity."}, status=status.HTTP_403_FORBIDDEN)
        reason = request.data.get("reason", "")
        if not isinstance(reason, str) or not reason.strip():
            return Response({"reason": "A void reason is required."}, status=status.HTTP_400_BAD_REQUEST)

        activity.void_reason = reason.strip()
        activity.voided_at = timezone.now()
        activity.voided_by = request.user
        activity.save(update_fields=["void_reason", "voided_at", "voided_by"])
        history(record, request.user, "activity_voided", {
            "activity_id": activity.id,
            "reason": activity.void_reason,
            "voided_at": activity.voided_at.isoformat(),
        })
        audit(request.user, "activity.voided", activity, after={"reason": activity.void_reason})
        return Response(ActivitySerializer(activity, context={"request": request}).data)

    @action(detail=True, methods=["get"])
    def history(self, request, pk=None):
        record = self.get_record(request, pk)
        if not can_read_record_history(request.user, record):
            return Response(status=status.HTTP_403_FORBIDDEN)
        return Response(HistorySerializer(record.history.all(), many=True).data)


class TaskViewSet(viewsets.ModelViewSet):
    serializer_class = TaskSerializer
    def get_queryset(self):
        query = Task.objects.filter(Q(assignee=self.request.user) | Q(record__in=visible_records(self.request.user))).select_related("assignee", "creator", "record", "deal")
        view = self.request.query_params.get("view")
        now = timezone.now()
        if view == "completed": query = query.filter(completed_at__isnull=False)
        elif view == "today": query = query.filter(completed_at__isnull=True, due_at__date=now.date())
        elif view == "overdue": query = query.filter(completed_at__isnull=True, due_at__lt=now)
        elif view == "upcoming": query = query.filter(completed_at__isnull=True, due_at__gt=now)
        return query
    def create(self, request, *args, **kwargs):
        if not request.user.has_crm_permission("create_task"): return Response(status=403)
        serializer = self.get_serializer(data=request.data); serializer.is_valid(raise_exception=True)
        record = serializer.validated_data.get("record")
        if record: assert_visible(request.user, record)
        task = serializer.save(creator=request.user, assignee=serializer.validated_data.get("assignee") or request.user)
        audit(request.user, "task.created", task, after={"title": task.title})
        return Response(self.get_serializer(task).data, status=201)
    @action(detail=True, methods=["post"])
    def complete(self, request, pk=None):
        task = self.get_object()
        if task.assignee_id != request.user.id and not request.user.has_crm_permission("complete_others_task"): return Response(status=403)
        task.completed_at = None if task.completed_at else timezone.now(); task.completed_by = request.user if task.completed_at else None; task.save(update_fields=["completed_at", "completed_by"])
        audit(request.user, "task.completed" if task.completed_at else "task.reopened", task)
        return Response(self.get_serializer(task).data)


def _client_ip(request):
    forwarded = request.META.get("HTTP_X_FORWARDED_FOR")
    return forwarded.split(",")[0].strip() if forwarded else request.META.get("REMOTE_ADDR")


def _auth_response(user, request):
    return {
        "id": user.id,
        "name": user.get_full_name() or user.username,
        "email": user.email,
        "mobile": user.mobile,
        "role": user.role.name if user.role else None,
        "permissions": user.role.permissions if user.role else [],
        "must_change_password": user.must_change_password,
    }


def _revoke_sessions(user, keep_key=None):
    qs = CrmSession.objects.filter(user=user, revoked_at__isnull=True)
    if keep_key:
        qs = qs.exclude(session_key=keep_key)
    keys = list(qs.values_list("session_key", flat=True))
    qs.update(revoked_at=timezone.now())
    if keys:
        Session.objects.filter(session_key__in=keys).delete()


@api_view(["POST"])
@permission_classes([AllowAny])
def login_view(request):
    identifier = normalize_identifier(request.data.get("username", ""))
    password = request.data.get("password", "")
    user = find_user(identifier)
    now = timezone.now()
    lockout_minutes = settings.LOGIN_LOCKOUT_MINUTES
    valid = bool(user and user.check_password(password))

    
    
    if user and user.locked_until and user.locked_until > now:
            audit(None,"auth.login.failed",user,after={"reason": "account_locked","ip": _client_ip(request),"browser": request.META.get("HTTP_USER_AGENT", "", ),  },) 
            return Response(
        {
            "detail": "Your account is temporarily locked due to multiple failed login attempts. Please try again after 15 minutes."
        },
        status=423,
    )
            
    if not valid or not user.is_active:
        if user and user.is_active:
            user.failed_login_attempts += 1
            if user.failed_login_attempts >= 5:
                user.failed_login_attempts = 0
                user.locked_until = now + timedelta(minutes=lockout_minutes)
                user.save(update_fields=["failed_login_attempts", "locked_until"])
                for admin_user in User.objects.select_related("role").filter(is_active=True):
                    if admin_user.has_crm_permission("manage_users"):
                        Notification.objects.create(user=admin_user, event_type="security.account_locked", message=f"Account locked after repeated sign-in failures: {user.email or user.username}.")
                audit(user, "auth.account_locked", user, after={"locked_until": user.locked_until.isoformat(), "ip": _client_ip(request), "browser": request.META.get("HTTP_USER_AGENT", "")})
            else:
                user.save(update_fields=["failed_login_attempts"])
        audit(user, "auth.login.failed", user, after={"reason": "invalid_credentials", "ip": _client_ip(request), "browser": request.META.get("HTTP_USER_AGENT", "")})
        return Response({"detail": "Your email/mobile number or password is incorrect."}, status=401)

    if user.locked_until and user.locked_until <= now:
        user.locked_until = None
    user.failed_login_attempts = 0
    user.save(update_fields=["failed_login_attempts", "locked_until"])
    login(request, user)
    CrmSession.objects.update_or_create(
        session_key=request.session.session_key,
        defaults={"user": user, "ip_address": _client_ip(request), "browser": request.META.get("HTTP_USER_AGENT", ""), "revoked_at": None},
    )
    audit(user, "auth.login.success", user, after={"ip": _client_ip(request), "browser": request.META.get("HTTP_USER_AGENT", "")})
    return Response(_auth_response(user, request))


@ensure_csrf_cookie
@api_view(["GET"])
@permission_classes([AllowAny])
def csrf_view(request):
    return Response(status=204)


@api_view(["POST"])
def logout_view(request):
    if request.session.session_key:
        CrmSession.objects.filter(session_key=request.session.session_key).update(revoked_at=timezone.now())
    audit(request.user, "auth.logout", request.user)
    logout(request)
    return Response(status=204)


@api_view(["GET"])
@permission_classes([AllowAny])
def me_view(request):
    if not request.user or not request.user.is_authenticated:
        return Response({"detail": "Authentication required."}, status=401)
    session = CrmSession.objects.filter(session_key=request.session.session_key, user=request.user, revoked_at__isnull=True).first()
    now = timezone.now()
    if session and now - session.last_seen_at >= timedelta(minutes=60):
        session.revoked_at = now
        session.save(update_fields=["revoked_at"])
        logout(request)
        return Response({"detail": "Authentication required."}, status=401)
    if not session or now - session.created_at >= timedelta(hours=8):
        if session:
            session.revoked_at = now
            session.save(update_fields=["revoked_at"])
        logout(request)
        return Response({"detail": "Authentication required."}, status=401)
    return Response(_auth_response(request.user, request))


@api_view(["POST"])
def password_change_view(request):
    current = request.data.get("current_password", "")
    new_password = request.data.get("new_password", "")
    if not request.user.check_password(current):
        return Response({"detail": "The current password is incorrect."}, status=400)
    try:
        validate_password(new_password, request.user)
    except Exception as exc:
        return Response({"detail": "; ".join(exc.messages)}, status=400)
    if not any(c.isupper() for c in new_password) or not any(not c.isalnum() for c in new_password):
        return Response({"detail": "Password must contain an uppercase letter and a special character."}, status=400)
    request.user.set_password(new_password)
    request.user.must_change_password = False
    request.user.save(update_fields=["password", "must_change_password"])
    current_key = request.session.session_key
    _revoke_sessions(request.user, keep_key=current_key)
    request.session.cycle_key()
    CrmSession.objects.filter(session_key=current_key).update(session_key=request.session.session_key, last_seen_at=timezone.now())
    audit(request.user, "auth.password.changed", request.user)
    return Response({"detail": "Password changed successfully."})


@api_view(["POST"])
@permission_classes([AllowAny])
def password_reset_request_view(request):
    # Forgot Password accepts EMAIL ONLY.
    # Login continues to support both email and mobile through find_user().
    email = str(request.data.get("identifier", "")).strip().lower()
    generic = {"detail": "If an account matches, a one-time code has been sent."}

    # Do not accept mobile numbers or other identifiers for password reset.
    if not email or "@" not in email:
        return Response(generic)

    user = User.objects.filter(email__iexact=email, is_active=True).first()
    if not user:
        return Response(generic)

    PasswordResetToken.objects.filter(
        user=user,
        used_at__isnull=True,
    ).update(used_at=timezone.now())

    code = f"{secrets.randbelow(1_000_000):06d}"
    PasswordResetToken.objects.create(
        user=user,
        code_hash=hashlib.sha256(code.encode()).hexdigest(),
        expires_at=timezone.now() + timedelta(minutes=10),
    )

    send_mail(
        "LiteCRM password reset code",
        f"Your LiteCRM password reset code is {code}. It expires in 10 minutes.",
        None,
        [user.email],
        fail_silently=True,
    )
    audit(user, "auth.password.reset_requested", user)
    return Response(generic)

@api_view(["POST"])
@permission_classes([AllowAny])
def password_reset_verify_view(request):
    identifier = normalize_identifier(
        request.data.get("identifier", "")
    )
    code = str(request.data.get("code", "")).strip()

    user = find_user(identifier)

    if not user or not user.is_active:
        return Response(
            {"detail": "Invalid OTP."},
            status=400,
        )

    token = (
        user.password_reset_tokens
        .filter(
            used_at__isnull=True,
            expires_at__gt=timezone.now(),
        )
        .order_by("-created_at")
        .first()
    )

    if not token:
        return Response(
            {"detail": "Invalid OTP or OTP has expired."},
            status=400,
        )

    if not secrets.compare_digest(
        token.code_hash,
        hashlib.sha256(code.encode()).hexdigest(),
    ):
        return Response(
            {"detail": "Invalid OTP."},
            status=400,
        )

    return Response(
        {"detail": "OTP verified successfully."},
        status=200,
    )


@api_view(["POST"])
@permission_classes([AllowAny])
def password_reset_confirm_view(request):
    # Password reset confirmation also accepts EMAIL ONLY.
    email = str(request.data.get("identifier", "")).strip().lower()
    code = request.data.get("code", "")
    new_password = request.data.get("new_password", "")

    if not email or "@" not in email:
        return Response({"detail": "The code is invalid or expired."}, status=400)

    user = User.objects.filter(email__iexact=email, is_active=True).first()
    if not user:
        return Response({"detail": "The code is invalid or expired."}, status=400)
    token = user.password_reset_tokens.filter(used_at__isnull=True, expires_at__gt=timezone.now()).order_by("-created_at").first()
    if not token or not secrets.compare_digest(token.code_hash, hashlib.sha256(code.encode()).hexdigest()):
        return Response({"detail": "The code is invalid or expired."}, status=400)
    try:
        validate_password(new_password, user)
    except Exception as exc:
        return Response({"detail": "; ".join(exc.messages)}, status=400)
    if not any(c.isupper() for c in new_password) or not any(not c.isalnum() for c in new_password):
        return Response({"detail": "Password must contain an uppercase letter and a special character."}, status=400)
    user.set_password(new_password)
    user.must_change_password = False
    user.save(update_fields=["password", "must_change_password"])
    token.used_at = timezone.now(); token.save(update_fields=["used_at"])
    _revoke_sessions(user)
    audit(user, "auth.password.reset_completed", user)
    return Response({"detail": "Password reset successfully."})


@api_view(["GET"])
def sessions_view(request):
    current = request.session.session_key
    rows = CrmSession.objects.filter(user=request.user, revoked_at__isnull=True).order_by("-last_seen_at")
    return Response([{"id": row.id, "browser": row.browser, "ip_address": row.ip_address, "created_at": row.created_at, "last_seen_at": row.last_seen_at, "current": row.session_key == current} for row in rows])


@api_view(["POST"])
def sessions_revoke_all_view(request):
    audit(request.user, "auth.sessions.revoked_all", request.user)
    _revoke_sessions(request.user)
    logout(request)
    return Response(status=204)


@api_view(["POST"])
def admin_revoke_user_sessions_view(request, user_id):
    if not request.user.has_crm_permission("manage_users"):
        return Response(status=403)
    target = User.objects.filter(pk=user_id).first()
    if not target:
        return Response(status=404)
    _revoke_sessions(target)
    audit(request.user, "auth.sessions.admin_revoked", target)
    return Response({"detail": "User sessions ended."})


@api_view(["POST"])
def admin_reset_password_view(request, user_id):
    if not request.user.has_crm_permission("manage_users"):
        return Response(status=403)
    target = User.objects.filter(pk=user_id).first()
    if not target:
        return Response(status=404)
    new_password = request.data.get("new_password", "")
    try:
        validate_password(new_password, target)
    except Exception as exc:
        return Response({"detail": "; ".join(exc.messages)}, status=400)
    if not any(c.isupper() for c in new_password) or not any(not c.isalnum() for c in new_password):
        return Response({"detail": "Password must contain an uppercase letter and a special character."}, status=400)
    target.set_password(new_password); target.must_change_password = True; target.save(update_fields=["password", "must_change_password"])
    _revoke_sessions(target)
    audit(request.user, "auth.password.admin_reset", target)
    return Response({"detail": "Password reset. The user must change it at next sign-in."})